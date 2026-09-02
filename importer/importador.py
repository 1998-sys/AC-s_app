# ----------------------------------------------------------------
# Project name  : AC's Generator (CertiFlow)
# Module        : importer.importador
# Created       : 04-05-2026
# Programmer(s) : Matheus Bandeira
# ----------------------------------------------------------------
# Remarks       : Reads an xlsx spreadsheet, categorizes each row against the instrument registry, and persists the accepted rows to the database.
#                 Lê uma planilha xlsx, categoriza cada linha em relação ao cadastro de instrumentos e grava no banco as linhas aceitas.
# ----------------------------------------------------------------
# Copyright (c) ODS Metering Systems
# ----------------------------------------------------------------

import openpyxl
from collections import defaultdict
from importer.validador import validar_linha
from data.utils_db import (
    inserir_instrumento, inserir_placa,
    atualizar_sn, atualizar_sn_sensor, atualizar_range, atualizar_sn_placa,
    atualizar_campos_extras,
)

COLUNAS_OBRIGATORIAS = {"tag", "sn_instrumento", "tipo"}

_SUFIXOS_PRESSAO = {"PT", "DPT", "PDT"}
_SUFIXOS_TEMP    = {"TT"}


def _segmentos(tag):
    """Returns the set of TAG segments (separated by "-"), in uppercase."""
    return {s.upper() for s in tag.split("-")}


def _e_grupo_mvs(tags):
    """Indicates whether the set of TAGs corresponds to an MVS (both pressure and temperature suffixes present)."""
    todos = set()
    for t in tags:
        todos |= _segmentos(t)
    return bool(todos & _SUFIXOS_PRESSAO) and bool(todos & _SUFIXOS_TEMP)


def _resolver_mvs(resultado):
    """Promotes to "inserir" (insert) the mvs_candidato groups whose TAG suffix pattern indicates an MVS.

    Groups the items in `resultado["mvs_candidato"]` by serial number (NS) and, when the
    group's set of TAGs indicates pressure + temperature under the same NS, marks the type
    as "MVS" and moves the items to `resultado["inserir"]`. The remaining groups stay in
    mvs_candidato for manual user confirmation.

    Args:
        resultado: categorization dict (same format returned by `ler_xlsx`), mutated in
            place — `mvs_candidato` and `inserir` are updated directly on the dict.
    """
    grupos = defaultdict(list)
    for item in resultado["mvs_candidato"]:
        grupos[item["sn"]].append(item)

    nao_auto = []
    for sn, itens in grupos.items():
        todas_tags = {i["tag"] for i in itens} | {i["tag_existente"] for i in itens}
        if _e_grupo_mvs(todas_tags):
            for item in itens:
                item["tipo"] = "MVS"
                resultado["inserir"].append(item)
        else:
            nao_auto.extend(itens)

    resultado["mvs_candidato"] = nao_auto


def ler_xlsx(caminho_xlsx):
    """Reads the xlsx spreadsheet and categorizes each row against the instrument registry, without writing to the database.

    Args:
        caminho_xlsx: path of the xlsx file to be read.

    Returns:
        dict with the keys "inserir", "mantido", "divergente", "bloqueado", "aviso" and
        "mvs_candidato", each containing the list of rows classified under that category.

    Notes:
        Data is read starting at row 3 (rows 1 and 2 are header/description).
        Fully empty rows are ignored.
    """
    try:
        wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
        ws = wb.active
    except Exception as e:
        raise ValueError(f"Erro ao abrir o arquivo: {e}")

    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    faltando = COLUNAS_OBRIGATORIAS - set(headers)
    if faltando:
        raise ValueError(f"Colunas obrigatórias ausentes no xlsx: {', '.join(sorted(faltando))}")

    idx = {h: i for i, h in enumerate(headers)}

    def get(row, col):
        """Reads the cell value from `row` in column `col` (by header name), as a trimmed string, or None if absent."""
        i = idx.get(col)
        if i is None:
            return None
        v = row[i].value
        return str(v).strip() if v is not None else None

    resultado = {"inserir": [], "mantido": [], "divergente": [], "bloqueado": [], "aviso": [], "mvs_candidato": []}

    for n, row in enumerate(ws.iter_rows(min_row=3), start=3):
        if all(c.value is None for c in row):
            continue

        tag  = (get(row, "tag") or "").upper()
        sn   = get(row, "sn_instrumento")
        tipo = (get(row, "tipo") or "").upper()
        sn_sensor = get(row, "sn_sensor")
        min_raw   = get(row, "min_range")
        max_raw   = get(row, "max_range")
        sistema   = get(row, "sistema")
        aplicacao = get(row, "aplicação") or get(row, "aplicacao")
        ativo     = get(row, "ativo")

        categoria, payload = validar_linha(
            n, tag, sn, tipo, sn_sensor, min_raw, max_raw,
            sistema, aplicacao, ativo,
        )
        resultado[categoria].append(payload)

    _resolver_mvs(resultado)
    return resultado


def executar(resultado, sobrescrever):
    """Persists to the database the records classified as new or confirmed for overwrite.

    Args:
        resultado: dict returned by `ler_xlsx`; the items in `resultado["inserir"]` are
            inserted as new records.
        sobrescrever: list of items from `resultado["divergente"]` that the user confirmed
            to overwrite (updates SN, sensor SN, range, and extra fields of the existing record).
    """
    for item in resultado["inserir"]:
        if item["tipo"] == "PO":
            inserir_placa(
                item["tag"], item["sn"],
                item.get("sistema"), item.get("aplicacao"), item.get("ativo"),
            )
        else:
            inserir_instrumento(
                item["tag"], item["sn"],
                item.get("sn_sensor"),
                item.get("min_range"),
                item.get("max_range"),
                item.get("sistema"),
                item.get("aplicacao"),
                item.get("ativo"),
                tipo=item["tipo"],
            )

    for item in sobrescrever:
        if item["tipo"] == "PO":
            atualizar_sn_placa(item["tag"], item["sn"])
        else:
            atualizar_sn(item["tag"], item["sn"])
            if item.get("sn_sensor") is not None:
                atualizar_sn_sensor(item["tag"], item["sn_sensor"])
            if item.get("min_range") is not None and item.get("max_range") is not None:
                atualizar_range(item["tag"], item["min_range"], item["max_range"])
        atualizar_campos_extras(
            item["tag"],
            item.get("sistema"), item.get("aplicacao"), item.get("ativo"),
        )
