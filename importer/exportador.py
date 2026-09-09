# ----------------------------------------------------------------
# Project name  : AC's Generator (CertiFlow)
# Module        : importer.exportador
# Created       : 15-05-2026
# Programmer(s) : Matheus Bandeira
# ----------------------------------------------------------------
# Remarks       : Exports the full instrument registry to an xlsx file formatted like the import template.
# ----------------------------------------------------------------
# Copyright (c) ODS Metering Systems
# ----------------------------------------------------------------

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from data.utils_db import listar_todos
from data.utils_fs import pasta_documentos

# Colors identical to the import template
_FILL_HEADER = PatternFill("solid", fgColor="C0392B")   # ODS red
_FILL_DESC   = PatternFill("solid", fgColor="F2F3F4")   # light gray
_FILL_EVEN   = PatternFill("solid", fgColor="F8F9FA")   # very light gray (even rows)

_FONT_HEADER = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_FONT_NORMAL = Font(name="Calibri", size=10)

_BORDER_THIN = Border(
    left=Side(style="thin", color="DEE2E6"),
    right=Side(style="thin", color="DEE2E6"),
    top=Side(style="thin", color="DEE2E6"),
    bottom=Side(style="thin", color="DEE2E6"),
)

_COLUNAS = [
    ("tag",            "Obrigatório",               18),
    ("sn_instrumento", "Obrigatório",               22),
    ("tipo",           "SEC / PO / MVS",            12),
    ("sn_sensor",      "Opcional — só SEC/MVS",     22),
    ("min_range",      "Opcional — só SEC/MVS",     14),
    ("max_range",      "Opcional — só SEC/MVS",     14),
    ("sistema",        "Opcional",                  20),
    ("aplicacao",      "Opcional",                  20),
    ("ativo",          "Opcional",                  16),
]


def exportar() -> str:
    """Exports every instrument registered in the database to an xlsx file, in the same format as the import template.

    Returns:
        str: absolute path of the generated xlsx file.
    """
    instrumentos = listar_todos()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "instrumentos"

    # ── Row 1: headers ───────────────────────────────────────────────────
    for col_idx, (campo, _, largura) in enumerate(_COLUNAS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=campo)
        cell.font      = _FONT_HEADER
        cell.fill      = _FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _BORDER_THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    ws.row_dimensions[1].height = 22

    # ── Row 2: descriptions ──────────────────────────────────────────────
    for col_idx, (_, descricao, _) in enumerate(_COLUNAS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=descricao)
        cell.font      = Font(italic=True, color="555555", name="Calibri", size=9)
        cell.fill      = _FILL_DESC
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _BORDER_THIN

    ws.row_dimensions[2].height = 16

    # ── Data rows ─────────────────────────────────────────────────────────
    for data_idx, inst in enumerate(instrumentos):
        row_num = data_idx + 3
        fill    = _FILL_EVEN if data_idx % 2 == 1 else None
        valores = [
            inst["tag"],
            inst["sn_instrumento"],
            inst["tipo"],
            inst["sn_sensor"],
            inst["min_range"],
            inst["max_range"],
            inst["sistema"],
            inst["aplicacao"],
            inst["ativo"],
        ]
        for col_idx, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=valor)
            cell.font   = _FONT_NORMAL
            cell.border = _BORDER_THIN
            cell.alignment = Alignment(vertical="center")
            if fill:
                cell.fill = fill

        ws.row_dimensions[row_num].height = 15

    # ── Freeze header ────────────────────────────────────────────────────
    ws.freeze_panes = "A3"

    # ── Automatic filter on headers ──────────────────────────────────────
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUNAS))}1"

    # ── Save ─────────────────────────────────────────────────────────────
    stamp   = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = str(pasta_documentos("AC's Generator/Exportações") / f"instrumentos_{stamp}.xlsx")
    wb.save(caminho)
    return caminho
