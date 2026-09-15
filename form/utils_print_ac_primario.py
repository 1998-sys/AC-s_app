# ----------------------------------------------------------------
# Project name  : AC's Generator (CertiFlow)
# Module        : form.utils_print_ac_primario
# Created       : 15-09-2026
# Programmer(s) : Matheus Bandeira
# ----------------------------------------------------------------
# Remarks       : Fills the primary flow meter Critical Analysis (AC) Excel template comparing the current calibration against the previous one, and exports both XLSX and PDF.
# ----------------------------------------------------------------
# Copyright (c) ODS Metering Systems
# ----------------------------------------------------------------

import os
from datetime import datetime

import openpyxl
import win32com.client as win32

from form.utils_print_linearizacao import parear_pontos_anterior

SHEET = "Sheet1"

# ── Header (rows 6-14) ──────────────────────────────────────────────────
CELL_TAG = "C6"
CELL_INSTRUMENTO = "C7"
CELL_FABRICANTE = "C8"
CELL_MODELO = "C9"
CELL_NUM_SERIE = "C10"
CELL_FAIXA_MEDICAO = "C11"
CELL_PRESSAO_CALIBRACAO = "C12"
CELL_DN = "C13"
CELL_UNIDADE_CERTIFICADO = "C14"

CELL_TIPO_EQUIPAMENTO = "H6"
CELL_CERTIFICADO_ANTERIOR = "H7"
CELL_LABORATORIO_ANTERIOR = "H8"
CELL_DATA_CAL_ANTERIOR = "H9"
CELL_CERTIFICADO_ATUAL = "H10"
CELL_LABORATORIO_ATUAL = "H11"
CELL_DATA_CAL_ATUAL = "H12"
CELL_VALIDADE_PADRAO = "H13"

# Fixed values (see tasks/TAREFAS_ac_medidor_primario_prio.md — assumed
# from the template's own example, not present as a distinct field in the
# flow meter XML schema).
INSTRUMENTO_PADRAO = "Medidor de Vazão"
UNIDADE_CERTIFICADO_PADRAO = "SI"

# ── Acceptance criteria (rows 17-19) ─────────────────────────────────────
CELL_ERRO_MAXIMO = "D17"
CELL_REPETIBILIDADE_LIMITE = "D18"
CELL_VALIDACAO_MF_LIMITE = "D19"

# Confirmed with the user (2026-09-15) — percentage points, matching the
# cells' own number format ("0.00", not "0.00%": 0.2 means "0,2%").
CRITERIOS_ACEITACAO = {
    "Master Meter": {
        "erro_maximo": 0.2,
        "repetibilidade": 0.02,
        "validacao_mf": None,  # does not apply
    },
    "Fiscal": {
        "erro_maximo": 0.2,
        "repetibilidade": 0.05,
        "validacao_mf": 0.25,
    },
    "Transferência de Custódia": {
        "erro_maximo": 0.2,
        "repetibilidade": 0.05,
        "validacao_mf": 0.25,
    },
    "Apropriação": {
        "erro_maximo": 0.6,
        "repetibilidade": 0.4,
        "validacao_mf": 2.0,
    },
    "Operacional": {
        "erro_maximo": 0.6,
        "repetibilidade": 0.4,
        "validacao_mf": 2.0,
    },
}

# ── Calibration point tables — up to 10 points each ─────────────────────
TABELA_ERRO_LINHA_INI = 28
TABELA_ERRO_LINHA_FIM_MAX = 37
TABELA_MF_LINHA_INI = 56
TABELA_MF_LINHA_FIM_MAX = 65
TABELA_REPETIBILIDADE_LINHA_INI = 86
TABELA_REPETIBILIDADE_LINHA_FIM_MAX = 95

CELL_HISTORICO_MF = "C53"          # "Possui histórico?" dropdown (SIM/NÃO)
CELL_HISTORICO_REPETIBILIDADE = "C83"

CELL_DATA_ANALISE = "D114"


def _data_iso_para_data(data_iso: str):
    """Converts an ISO date string (YYYY-MM-DD) to a `datetime` object for a date-typed cell.

    Args:
        data_iso: Date in ISO format, e.g. "2025-09-26".

    Returns:
        datetime.datetime: parsed date, or None if `data_iso` is empty or
        cannot be parsed (leaves the cell untouched rather than writing a
        wrong value).
    """
    if not data_iso:
        return None
    try:
        return datetime.strptime(data_iso, "%Y-%m-%d")
    except ValueError:
        return None


def _validade_padrao_texto(padroes: list) -> str:
    """Builds the multi-line "Validade Padrão" text from the certificate's reference standards.

    Args:
        padroes: list of {"identificador", "validade"} dicts (see
            `xml_extractor_FT.extrair_dados_ft`'s "padroes" key) — one per
            reference standard used in the calibration.

    Returns:
        str: one "<identificador>: <validade DD/MM/YY>" line per standard,
        joined with "\\n" (matches the template's own example format,
        e.g. "Padrão 16 HV: 23/01/26\\nPadrão 15 HV: 23/03/25"); empty
        string if there are no standards.
    """
    linhas = []
    for p in padroes or []:
        data = _data_iso_para_data(p.get("validade"))
        validade_fmt = data.strftime("%d/%m/%y") if data else (p.get("validade") or "")
        identificador = p.get("identificador") or ""
        if identificador or validade_fmt:
            linhas.append(f"{identificador}: {validade_fmt}".strip(": "))
    return "\n".join(linhas)


def gerar_ac_primario(dados_atual: dict, dados_anterior: dict, aplicacao: str, caminho_xml: str) -> str:
    """Fills Template_AC_PRIM_PRIO.xlsx comparing the current calibration against the previous one, and exports both XLSX and PDF.

    Requires a previous certificate to compare against — same requirement
    already in place for the Presumed Failure report (see
    `form.utils_print_linearizacao.gerar_falha_presumida`); both are
    generated from the same pair of XMLs, so the user is only asked for
    the previous certificate once.

    Args:
        dados_atual: current certificate's data (`xml_extractor_FT.extrair_dados_ft`
            output on the current XML).
        dados_anterior: previous certificate's data (same function's output
            on the previous XML).
        aplicacao: one of `CRITERIOS_ACEITACAO`'s keys ("Master Meter",
            "Fiscal", "Transferência de Custódia", "Apropriação",
            "Operacional") — resolves the 3 acceptance thresholds written
            to D17/D18/D19.
        caminho_xml: path of the current XML; used only to determine the
            output folder (same folder as the Linearização/Falha
            Presumida files).

    Returns:
        str: path of the generated PDF (the XLSX is saved alongside it,
        same base name).

    Raises:
        KeyError: if `aplicacao` isn't one of `CRITERIOS_ACEITACAO`'s keys.
        ValueError: if the certificate has more calibration points than the
            template supports (10), or if `dados_anterior` has no points to
            pair against (see `parear_pontos_anterior`).
        PermissionError: if the output PDF already exists and is open.
    """
    criterios = CRITERIOS_ACEITACAO[aplicacao]

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    caminho_template = os.path.join(base_dir, "templates", "Template_AC_PRIM_PRIO.xlsx")

    wb = openpyxl.load_workbook(caminho_template)
    ws = wb[SHEET]

    # ── Header ────────────────────────────────────────────────────────
    ws[CELL_TAG] = dados_atual.get("tag", "")
    ws[CELL_INSTRUMENTO] = INSTRUMENTO_PADRAO
    ws[CELL_FABRICANTE] = dados_atual.get("fabricante", "")
    ws[CELL_MODELO] = dados_atual.get("modelo", "")
    ws[CELL_NUM_SERIE] = dados_atual.get("num_serie", "")
    ws[CELL_FAIXA_MEDICAO] = dados_atual.get("faixa_calibrada", "")
    if dados_atual.get("pressao_calibracao"):
        unidade = dados_atual.get("pressao_calibracao_unidade", "")
        ws[CELL_PRESSAO_CALIBRACAO] = f"{dados_atual['pressao_calibracao']} {unidade}".strip()
    ws[CELL_DN] = dados_atual.get("diametro", "")
    ws[CELL_UNIDADE_CERTIFICADO] = UNIDADE_CERTIFICADO_PADRAO

    ws[CELL_TIPO_EQUIPAMENTO] = dados_atual.get("tipo", "")
    ws[CELL_CERTIFICADO_ANTERIOR] = dados_anterior.get("numero_certificado", "")
    ws[CELL_LABORATORIO_ANTERIOR] = dados_anterior.get("laboratorio", "")
    ws[CELL_DATA_CAL_ANTERIOR] = _data_iso_para_data(dados_anterior.get("data_calibracao"))
    ws[CELL_CERTIFICADO_ATUAL] = dados_atual.get("numero_certificado", "")
    ws[CELL_LABORATORIO_ATUAL] = dados_atual.get("laboratorio", "")
    ws[CELL_DATA_CAL_ATUAL] = _data_iso_para_data(dados_atual.get("data_calibracao"))
    ws[CELL_VALIDADE_PADRAO] = _validade_padrao_texto(dados_atual.get("padroes"))

    # ── Acceptance criteria ──────────────────────────────────────────
    ws[CELL_ERRO_MAXIMO] = criterios["erro_maximo"]
    ws[CELL_REPETIBILIDADE_LIMITE] = criterios["repetibilidade"]
    ws[CELL_VALIDACAO_MF_LIMITE] = criterios["validacao_mf"]

    # ── Point-by-point comparison tables ─────────────────────────────
    pontos_atual = dados_atual.get("pontos", [])
    max_pontos = TABELA_ERRO_LINHA_FIM_MAX - TABELA_ERRO_LINHA_INI + 1
    if len(pontos_atual) > max_pontos:
        raise ValueError(
            f"Template de AC de medidor primário suporta no máximo {max_pontos} "
            f"pontos de calibração (certificado tem {len(pontos_atual)})."
        )

    pontos_pareados = parear_pontos_anterior(pontos_atual, dados_anterior.get("pontos", []))

    for i, (p_atual, p_anterior) in enumerate(zip(pontos_atual, pontos_pareados)):
        linha_erro = TABELA_ERRO_LINHA_INI + i
        ws[f"C{linha_erro}"] = p_atual["vazao"]
        ws[f"D{linha_erro}"] = p_anterior["erro_pct"]
        ws[f"E{linha_erro}"] = p_atual["erro_pct"]
        ws[f"F{linha_erro}"] = p_anterior["incerteza"]
        ws[f"G{linha_erro}"] = p_atual["incerteza"]

        linha_mf = TABELA_MF_LINHA_INI + i
        ws[f"D{linha_mf}"] = p_anterior["meter_factor"]
        ws[f"E{linha_mf}"] = p_atual["meter_factor"]

        linha_rep = TABELA_REPETIBILIDADE_LINHA_INI + i
        ws[f"D{linha_rep}"] = p_anterior["repetibilidade"]
        ws[f"E{linha_rep}"] = p_atual["repetibilidade"]

    # The template ships with sample data pre-filled in ALL 10 rows of each
    # table (not blank placeholders) — a certificate with fewer points than
    # that would otherwise leave the template's own example values sitting
    # in the unused rows (they even leak into the "Análise Paramétrica"
    # charts, which plot straight off these cells). Clear every row past
    # the real point count, up to each table's max row.
    #
    # Also clears each table's "echo" column (H in Erro, H in Meter Factor,
    # not needed in Repetibilidade) even though it's a template formula, not
    # something we normally write to: those formulas don't check for an
    # empty row like the actual Resultado column does, so an untouched
    # unused row would show a stray repeated threshold value (Erro's H,
    # harmless) or literally "#DIV/0!" (Meter Factor's H = "=(E-D)/D",
    # divides by the now-blank D) — a real error visible in the exported
    # PDF, not just cosmetic.
    for linha in range(TABELA_ERRO_LINHA_INI + len(pontos_atual), TABELA_ERRO_LINHA_FIM_MAX + 1):
        for col in ("C", "D", "E", "F", "G", "H"):
            ws[f"{col}{linha}"] = None
    for linha in range(TABELA_MF_LINHA_INI + len(pontos_atual), TABELA_MF_LINHA_FIM_MAX + 1):
        for col in ("D", "E", "G", "H"):
            ws[f"{col}{linha}"] = None
    for linha in range(TABELA_REPETIBILIDADE_LINHA_INI + len(pontos_atual), TABELA_REPETIBILIDADE_LINHA_FIM_MAX + 1):
        for col in ("D", "E"):
            ws[f"{col}{linha}"] = None

    ws[CELL_HISTORICO_MF] = "SIM"
    ws[CELL_HISTORICO_REPETIBILIDADE] = "SIM"

    ws[CELL_DATA_ANALISE] = datetime.now()

    # ── Save output XLSX ──────────────────────────────────────────────
    pasta_saida = os.path.dirname(os.path.abspath(caminho_xml))
    cert_limpo = dados_atual.get("numero_certificado", "").replace(" ", "")
    tag_limpa = dados_atual.get("tag", "").replace(" ", "")
    nome_xlsx = f"{cert_limpo}_{tag_limpa}_AC_MEDIDOR_PRIMARIO.xlsx"
    caminho_xlsx = os.path.join(pasta_saida, nome_xlsx)

    wb.save(caminho_xlsx)

    # ── Export PDF via Excel COM ────────────────────────────────────────
    nome_pdf = f"{cert_limpo}_{tag_limpa}_AC_MEDIDOR_PRIMARIO.pdf"
    caminho_pdf = os.path.join(pasta_saida, nome_pdf)

    if os.path.exists(caminho_pdf):
        try:
            os.remove(caminho_pdf)
        except PermissionError:
            raise PermissionError(
                f"O arquivo PDF está aberto e não pode ser sobrescrito:\n{caminho_pdf}"
            )

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False
    excel.Interactive = False

    try:
        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_xlsx))
        excel.Calculate()
        ws_excel = wb_excel.Worksheets(SHEET)
        ws_excel.ExportAsFixedFormat(0, caminho_pdf)
        wb_excel.Close(SaveChanges=False)
    finally:
        excel.Quit()

    return caminho_pdf
