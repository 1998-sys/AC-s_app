# ----------------------------------------------------------------
# Project name  : AC's Generator (CertiFlow)
# Module        : form.po_templates
# Created       : 24-08-2026
# Programmer(s) : Matheus Bandeira
# ----------------------------------------------------------------
# Remarks       : Fills the orifice plate (PO) AC Excel templates for ORIGEM, YINSON and YINSON ATLANTA and exports them to PDF via Excel COM automation.
#                 Preenche os templates Excel de AC de placa de orifício (PO) para ORIGEM, YINSON e YINSON ATLANTA e os exporta para PDF via automação COM do Excel.
# ----------------------------------------------------------------
# Copyright (c) ODS Metering Systems
# ----------------------------------------------------------------

import openpyxl
import os
import re
import shutil
import tempfile
import time
import uuid
import win32com.client as win32


def limpar_nome_arquivo(texto):
    """Removes characters invalid for file names and replaces spaces with underscores.

    Args:
        texto: Input text; if empty/None, returns an empty string.

    Returns:
        str: Sanitized text, with the characters \\/*?:"<>| removed and spaces
        converted to underscores.
    """
    if not texto:
        return ""

    texto = re.sub(r'[\\/*?:"<>|]', "", texto)
    texto = re.sub(r"\s+", "_", texto)
    return texto.strip()


def gerar_ac_po(config, dados, caminho_pdf_original, excel=None):
    """Fills an orifice plate AC Excel template (PO variant) and exports it as PDF via Excel COM.

    ORIGEM_PO, YINSON_PO and YINSON_ATLANTA_PO all use this exact same flow —
    copy the template to a temporary file, write a handful of fields,
    open it via Excel COM to recalculate formulas and export the PDF — differing
    only in the template used and which fields go into which cells (`config`).

    Args:
        config (dict): 'template', 'prefixo_temp', 'campos' (cell -> key in
            `dados`), 'data_cell' (cell for "Data: {report_date}"),
            'nome_arquivo_campo' (key in `dados` used in the final PDF's file name).
        dados (dict): Certificate fields.
        caminho_pdf_original (str): Path of the source PDF, used to
            determine the output folder.

    Returns:
        str: Absolute path of the generated PDF.

    Raises:
        PermissionError: If the output PDF is already open.
    """
    caminho_template = config["template"]

    caminho_temp = os.path.join(
        tempfile.gettempdir(),
        f"{config['prefixo_temp']}{uuid.uuid4().hex}.xlsx"
    )
    shutil.copy(caminho_template, caminho_temp)

    wb = openpyxl.load_workbook(caminho_temp)
    ws = wb["Template Formulário"]

    for celula, campo in config["campos"].items():
        ws[celula] = dados.get(campo)

    ws[config["data_cell"]] = f'Data: {dados.get("report_date")}'

    wb.save(caminho_temp)
    wb.close()

    time.sleep(1)
    pasta_saida = os.path.dirname(os.path.abspath(caminho_pdf_original))

    nome_base = limpar_nome_arquivo(dados.get(config["nome_arquivo_campo"], ""))
    nome_pdf = f"{nome_base}_AC.pdf"
    caminho_pdf_final = os.path.join(pasta_saida, nome_pdf)

    if os.path.exists(caminho_pdf_final):
        try:
            os.remove(caminho_pdf_final)
        except PermissionError:
            raise PermissionError(
                f"⚠ O PDF está aberto:\n{caminho_pdf_final}"
            )

    excel_proprio = excel is None
    if excel_proprio:
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

    try:
        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_temp))

        time.sleep(1)

        excel.Calculate()
        wb_excel.ExportAsFixedFormat(
            0,
            os.path.abspath(caminho_pdf_final)
        )

        wb_excel.Close(False)

    finally:
        if excel_proprio:
            excel.Quit()
        if os.path.exists(caminho_temp):
            try:
                os.remove(caminho_temp)
            except Exception:
                pass

    return caminho_pdf_final


CONFIG_ORIGEM_PO = {
    "template": "TemplateAC_PO_ORIGEM.xlsx",
    "prefixo_temp": "temp_ac_origem_",
    "campos": {
        "A6": "sn_inst",
        "A13": "certificado",
        "D13": "data_calibracao",
        "C6": "localizacao",
        "F6": "n_ac",
    },
    "data_cell": "D57",
    "nome_arquivo_campo": "n_ac",
}

CONFIG_YINSON_PO = {
    "template": "TemplateAC_PO_YINSON.xlsx",
    "prefixo_temp": "temp_ac_yinson_",
    "campos": {
        "A13": "certificado",
        "D13": "data_calibracao",
        "C6": "local",
        "F6": "sn_inst",
    },
    "data_cell": "D56",
    "nome_arquivo_campo": "certificado",
}

CONFIG_YINSON_ATLANTA_PO = {
    "template": "TemplateAC_PO_YINSON - ATLANTA.xlsx",
    "prefixo_temp": "temp_ac_yinson_atlanta_",
    "campos": {
        "A13": "certificado",
        "D13": "data_calibracao",
        "C6": "local",
        "F6": "sn_inst",
    },
    "data_cell": "D56",
    "nome_arquivo_campo": "certificado",
}


def gerar_ac_origem_PO(dados, caminho_pdf_original, excel=None):
    """Fills the ORIGEM PO AC Excel template and exports it as PDF via Excel COM."""
    return gerar_ac_po(CONFIG_ORIGEM_PO, dados, caminho_pdf_original, excel=excel)


def gerar_ac_yinson_PO(dados, caminho_pdf_original, excel=None):
    """Fills the YINSON PO AC Excel template and exports it as PDF via Excel COM."""
    return gerar_ac_po(CONFIG_YINSON_PO, dados, caminho_pdf_original, excel=excel)


def gerar_ac_yinson_atlanta_PO(dados, caminho_pdf_original, excel=None):
    """Fills the YINSON ATLANTA PO AC Excel template and exports it as PDF via Excel COM."""
    return gerar_ac_po(CONFIG_YINSON_ATLANTA_PO, dados, caminho_pdf_original, excel=excel)
