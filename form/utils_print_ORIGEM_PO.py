import openpyxl
import os
import win32com.client as win32
import tempfile
import shutil
import uuid
import time
import re


def limpar_nome_arquivo(texto):
    if not texto:
        return ""
    
    texto = re.sub(r'[\\/*?:"<>|]', "", texto)
    texto = re.sub(r"\s+", "_", texto)
    return texto.strip()


def gerar_ac_origem_PO(dados, caminho_pdf_original):

    caminho_template = "TemplateAC_PO_ORIGEM.xlsx"

    temp_dir = tempfile.gettempdir()
    caminho_temp = os.path.join(
        temp_dir,
        f"temp_ac_origem_{uuid.uuid4().hex}.xlsx"
    )

    shutil.copy(caminho_template, caminho_temp)

    wb = openpyxl.load_workbook(caminho_temp)
    ws = wb["Template Formulário"]

    ws["A6"] = dados.get("tag")
    ws["A13"] = dados.get("certificado")
    ws["D13"] = dados.get("data_calibracao")
    ws["D57"] = f'Data: {dados.get("report_date")}'
    ws["C6"] = dados.get("localizacao")
    ws["F6"] = dados.get('sap')

    wb.save(caminho_temp)
    wb.close()

    time.sleep(1)
    pasta_saida = os.path.dirname(os.path.abspath(caminho_pdf_original))

    certificado = limpar_nome_arquivo(dados.get("certificado", ""))
    tag_limpa = limpar_nome_arquivo(dados.get("tag", ""))

    nome_pdf = f"{certificado}_{tag_limpa}_AC.pdf"
    caminho_pdf_final = os.path.join(pasta_saida, nome_pdf)

    if os.path.exists(caminho_pdf_final):
        try:
            os.remove(caminho_pdf_final)
        except PermissionError:
            raise PermissionError(
                f"⚠ O PDF está aberto:\n{caminho_pdf_final}"
            )

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    try:
        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_temp))

        time.sleep(1)

        excel.Calculate()
        wb_excel.ExportAsFixedFormat(
            0,
            os.path.abspath(caminho_pdf_final)
        )

        wb_excel.Close(False)

    finally:
        excel.Quit()
        if os.path.exists(caminho_temp):
            try:
                os.remove(caminho_temp)
            except Exception:
                pass

    return caminho_pdf_final