# ----------------------------------------------------------------
# Project name  : AC's Generator (CertiFlow)
# Module        : form.utils_print_PRIO_PO
# Created       : 10-12-2025
# Programmer(s) : Matheus Bandeira
# ----------------------------------------------------------------
# Remarks       : Fills the AC PRIO orifice plate (PO) Excel template, adjusting borders and print area, and exports it to PDF via Excel COM automation.
#                 Preenche o template Excel de AC PRIO de placa de orifício (PO), ajustando bordas e área de impressão, e o exporta para PDF via automação COM do Excel.
# ----------------------------------------------------------------
# Copyright (c) ODS Metering Systems
# ----------------------------------------------------------------

import openpyxl
import win32com.client as win32
import os
import re
import shutil
import tempfile
import uuid
from openpyxl.utils import get_column_letter
from openpyxl.styles import Border, Side
from openpyxl.worksheet.page import PageMargins


def gerar_ac_prio_po(
    dados: dict,
    caminho_pdf_original,
    caminho_template = "templates/TemplateAC_PO_PRIO.xlsx",
    nome_aba = "Template Formulário",
    print_area_fixa: str | None = None,
    linhas_extra_topo: int = 2,
    aplicar_borda_fallback: bool = True,
    cor_faixa_teal_hex = "0099A8",
    excel=None
):
    """Fills the AC PRIO PO template and exports it as PDF via Excel COM.

    Applies borders, auto-detects the print area from cell content and
    overrides the page setup directly via COM to guarantee that "fit to
    page" is honored on export.

    Args:
        dados: Certificate fields.
        caminho_pdf_original: Path of the source PDF, used to determine the
            output folder.
        caminho_template: Path of the Excel template to fill.
        nome_aba: Name of the template sheet to use.
        print_area_fixa: Fixed print area (e.g. "A1:H50"); if None, it is
            computed automatically from the last filled cell.
        linhas_extra_topo: Extra rows added to the auto-detected print area,
            to give some slack below the content.
        aplicar_borda_fallback: If True, draws a thick bottom border right
            below the title "Análise crítica de calibração de placas
            de orifício", if it is found within the first 30 rows.
        cor_faixa_teal_hex: Color (hex, without "#") used for the fallback
            border.
        excel: Already-open Excel COM instance (reused across a batch). If
            None, opens and closes its own instance.

    Returns:
        str: Absolute path of the generated PDF.

    Raises:
        FileNotFoundError: If `caminho_template` does not exist.
        KeyError: If `nome_aba` does not exist in the template.
    """

    if not os.path.isfile(caminho_template):
        raise FileNotFoundError(f"Template não encontrado: {os.path.abspath(caminho_template)}")

    caminho_temp = os.path.join(
        tempfile.gettempdir(),
        f"temp_ac_prio_po_{uuid.uuid4().hex}.xlsx"
    )
    shutil.copy(caminho_template, caminho_temp)

    wb = openpyxl.load_workbook(caminho_temp)

    if nome_aba not in wb.sheetnames:
        raise KeyError(f"Aba '{nome_aba}' não encontrada.")

    ws = wb[nome_aba]

    ws["C5"] = dados.get("local", "")
    ws["C6"] = dados.get("local", "")
    ws["C7"] = dados.get("data_calibracao", "")
    ws["F5"] = dados.get("sn_inst", "")
    ws["F6"] = dados.get("certificado", "")
    ws["F7"] = "ODS Metering Systems"
    ws["G48"] = dados.get('report_date', "")

    if aplicar_borda_fallback:

        titulo = "Análise crítica de calibração de placas de orifício"
        titulo_row = None
        max_col = 1

        for row in ws.iter_rows(min_row=1, max_row=30):

            for cell in row:

                if cell.value not in (None, ""):
                    max_col = max(max_col, cell.column)

                    if isinstance(cell.value, str) and titulo.lower() in cell.value.lower():
                        titulo_row = cell.row

        if titulo_row:

            linha_borda = titulo_row + 1
            teal = cor_faixa_teal_hex.upper()

            border = Border(bottom=Side(style="thick", color=teal))

            for c in range(1, max_col + 1):

                addr = f"{get_column_letter(c)}{linha_borda}"
                ws[addr].border = ws[addr].border + border

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    ws.print_options.horizontalCentered = True

    ws.page_margins = PageMargins(
        top=0.3,
        bottom=0.3,
        left=0.3,
        right=0.3
    )

    if print_area_fixa is None:

        last_row = 1
        last_col = 1

        for row in ws.iter_rows():

            for cell in row:

                if cell.value not in (None, ""):

                    last_row = max(last_row, cell.row)
                    last_col = max(last_col, cell.column)

        last_row += linhas_extra_topo

        ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"

    else:

        ws.print_area = print_area_fixa


    wb.save(caminho_temp)
    wb.close()

    pasta_saida = os.path.dirname(os.path.abspath(caminho_pdf_original))

    def sanitize(txt):
        """Replaces characters invalid for a file/path name with a hyphen."""
        return re.sub(r'[<>:"/\\|?*]+', "-", str(txt or "")).strip()

    certificado = sanitize(dados.get("certificado")).replace(" ", "")
    tag = sanitize(dados.get("tag")).replace(" ", "")

    nome_pdf = f"{certificado}_{tag}_AC.pdf".strip("_")
    pdf_final = os.path.join(pasta_saida, nome_pdf)

    if os.path.exists(pdf_final):
        os.remove(pdf_final)


    excel_proprio = excel is None
    wb_excel = None

    try:

        if excel_proprio:
            excel = win32.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False

        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_temp))
        ws_excel = wb_excel.Worksheets(nome_aba)

        ps = ws_excel.PageSetup

        ps.Zoom = False
        ps.FitToPagesWide = 1
        ps.FitToPagesTall = 1

        ps.LeftMargin = excel.InchesToPoints(0.3)
        ps.RightMargin = excel.InchesToPoints(0.3)
        ps.TopMargin = excel.InchesToPoints(0.3)
        ps.BottomMargin = excel.InchesToPoints(0.3)

        ps.CenterHorizontally = True


        wb_excel.ExportAsFixedFormat(
            Type=0,
            Filename=pdf_final,
            Quality=0,
            IncludeDocProperties=True,
            IgnorePrintAreas=False,
            OpenAfterPublish=False
        )

    finally:

        try:
            if wb_excel:
                wb_excel.Close(SaveChanges=False)
        except:
            pass

        try:
            if excel_proprio and excel:
                excel.Quit()
        except:
            pass

        if os.path.exists(caminho_temp):
            try:
                os.remove(caminho_temp)
            except Exception:
                pass

    return pdf_final