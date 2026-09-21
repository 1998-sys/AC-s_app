# ----------------------------------------------------------------
# Project name  : AC's Generator (CertiFlow)
# Module        : form.utils_print_linearizacao
# Created       : 31-07-2026
# Programmer(s) : Matheus Bandeira
# ----------------------------------------------------------------
# Remarks       : Fills the flow meter linearization Excel template with the calibration points and exports both XLSX and PDF.
# ----------------------------------------------------------------
# Copyright (c) ODS Metering Systems
# ----------------------------------------------------------------

from copy import copy
import math
import openpyxl
from openpyxl.styles import Border, Side
import win32com.client as win32
import os
from datetime import datetime


# ── CELL MAP ─────────────────────────────────────────────────────────
# Addresses confirmed by opening Template_Linearizacao.xlsx (sheet "Linearização").
# Frequência (E), Status (S), the 2nd table "Dados a Serem Configurados" (rows
# 56-75), KF médio (L76) and the alarm limits (J79/M79) are already FORMULAS in
# the template itself — all derived from the columns we fill here
# (C/G/I/K/M/Q). Do not write to them: Excel recalculates them on its own
# when opening/exporting (including the pass/fail color, which is native
# conditional formatting in S22:T41 — no PatternFill needed).
# K-factor corrigido (O) is also a template formula, but the fixed
# ROUND()/format at 5 decimal places in the template only reaches 7
# significant figures when the value falls between 10 and 100 — for a
# typical K-Factor (close to 1), that only yields 6 digits. That's why this
# formula is rewritten per row in gerar_linearizacao (decimal places
# computed from the order of magnitude, same rule as the nominal K-Factor)
# instead of being left untouched like the others.
CELLS = {
    # Left header
    "cliente":         "D9",
    "instalacao":      "D10",
    "tag_sistema":     "D11",
    "aplicacao":       "D12",
    "sistema":         "D13",
    "data_calibracao": "D14",
    "tipo_medidor":    "D15",
    # Right header ("TAG" in M12 is formula "=D11", do not write)
    "num_certificado": "M9",
    "modelo":          "M10",
    "fabricante":      "M11",
    "num_serie":       "M13",
    "diametro":        "M14",
    "faixa_calibrada": "M15",
    # Nominal K-Factor
    "fator_k":         "D19",
    # KF médio (template formula "=AVERAGE(...)", in the 2nd table)
    "kf_medio":        "L76",
    # Calibration table (rows 22 to 41 — up to 20 points)
    "tabela_linha_ini": 22,
    "tabela_linha_fim": 41,
    "col_vazao":       "C",
    "col_vol_ref":     "G",
    "col_vol_med":     "I",
    "col_mf":          "K",
    "col_kfc":         "O",
    "col_erro":        "M",
    "col_incerteza":   "Q",
    # Column titles (the unit written in parentheses is dynamically
    # overwritten with the certificate's actual unit — see gerar_linearizacao).
    "titulo_vazao":     "C21",
    "titulo_vol_ref":   "G21",
    "titulo_vol_med":   "I21",
    # Signature block (default names already come filled in the template;
    # only overwritten if the user chooses to change them in the prompt).
    "elaborado_nome": "F96",
    "elaborado_data": "F98",
    "verificado_nome": "N96",
    # verificado_data (N98) is formula "=F98" in the template — do not overwrite.
}

# Columns with a border in the calibration table (B = outer left border to
# T = outer right border). Row 30 is used as the reference for the "normal"
# (thin) border in the middle of the table — see _ajustar_bordas_tabela.
_BORDA_COLS = ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T"]
_LINHA_BORDA_REF = 30

# The "Dados a Serem Configurados" table (row + TABELA2_OFFSET) is narrower
# than the one above — it only uses H to M (N°/Vazão/Frequência/K-factor
# corrigido); C to G and N to T are always left without a border there.
# Applying table 1's column range (B to T) to it leaves a thick border on
# columns N to T of the closing row, "leaking" outside table 2's actual box.
_BORDA_COLS_TABELA2 = ["H", "I", "J", "K", "L", "M"]

# Row offset from table 1 (calibration) to the 2nd table ("Dados a Serem
# Configurados") — e.g. row 22 of table 1 mirrors row 56 of table 2. Used
# both for hiding/showing rows and for the K-Factor Corrigido format
# (col. L in table 2, which mirrors col. O in table 1).
TABELA2_OFFSET = 34

# ── "Falha Presumida" sheet ──────────────────────────────────────────────
# Compares this certificate's meter factor against the previous calibration
# of the same meter (see gerar_falha_presumida). The sheet mirrors almost
# everything from "Linearização" via formula (header, calibration table,
# Diff MF / Fator de Correção / Status) — the only cells this module writes
# are the previous certificate's number and its meter factor per point.
SHEET_FALHA_PRESUMIDA = "Falha Presumida"
FP_CERT_ANTERIOR = "G85"          # previous certificate's number
FP_COL_MF_ANTERIOR = "G"          # previous certificate's meter factor, per row
FP_LINHA_INI = 87                 # mirrors Linearização's row 22 (offset 65)
FP_LINHA_FIM_MAX = 106            # mirrors Linearização's row 41 (up to 20 points)
FP_OFFSET_CALCULO = FP_LINHA_INI - CELLS["tabela_linha_ini"]  # 87 - 22 = 65
FP_LINHA_BORDA_REF_CALCULO = 90   # a "middle" row of the Cálculo Falha Presumida table
FP_BORDA_COLS_CALCULO = ["E", "G", "I", "K", "L", "N"]

# The calibration-table mirror (rows 22-41) has its own row visibility/border
# state on this sheet (Excel doesn't share hidden-row state between sheets),
# even though its cell values are 100% formula-mirrored from Linearização.
FP_MIRROR_LINHA_INI = 22
FP_MIRROR_LINHA_FIM_MAX = 41
FP_MIRROR_LINHA_BORDA_REF = 30
FP_MIRROR_BORDA_COLS = ["B", "C", "E", "G", "I", "K", "M", "O", "Q", "S", "T"]


def _data_br(data_iso: str) -> str:
    """Converts a date from ISO format (YYYY-MM-DD) to BR format (DD/MM/YYYY).

    Args:
        data_iso: Date in ISO format, e.g. "2026-04-25".

    Returns:
        str: Date in "DD/MM/YYYY" format, e.g. "25/04/2026"; returns
        `data_iso` unchanged if it cannot be parsed.
    """
    try:
        return datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        return data_iso


# Known clients recognized in the Linearização/Falha Presumida "Cliente"
# field — matched as a case-insensitive substring, so a raw XML value like
# "Prio Forte S/A" or a folder segment like ".../PRIO/..." both resolve to
# the same normalized "PRIO". Includes "ODS" itself: on some certificates
# CLIENTE/NOME names ODS as the calibration lab's own contracting client
# (e.g. "ODS do Brasil Sistemas de Medição LTDA") rather than the end oil
# company — a legitimate value for this field, not a detection error.
CLIENTES_CONHECIDOS = ("PRIO", "YINSON", "ORIGEM", "SBM", "PETROBRAS", "ODS")


def identificar_cliente(nome_cliente_xml: str, caminho_xml: str) -> str:
    """Suggests a default "Cliente" value from the certificate's own CLIENTE/NOME text, falling back to the file path.

    `gui.pdf_service` uses the result only as a pre-selected default for
    its "Cliente" prompt field (a fixed PRIO/YINSON/ORIGEM selection, since
    that value will drive AC template routing for primary meters — see the
    "Primary meter ACs" roadmap item) — if this returns something outside
    that set (e.g. "ODS", or a client only recognized via the broader
    `CLIENTES_CONHECIDOS` list), the field just starts unselected and the
    user picks one of the 3 explicitly. `gerar_linearizacao` also calls
    this directly as its own fallback when `dados["cliente"]` isn't already
    set (e.g. called standalone, outside that prompt flow) — there, any of
    `CLIENTES_CONHECIDOS` is written into the sheet as-is (free text cell,
    no routing implication yet).

    Args:
        nome_cliente_xml: raw text of the XML's CLIENTE/NOME element (see
            `xml_extractor_FT.extrair_dados_ft`'s "cliente_xml" key) —
            checked first since it doesn't depend on where the file happens
            to be saved on disk (a loose/dropped file saved to a generic
            folder has no client name in its path at all).
        caminho_xml: file path (e.g. ".../PRIO/..."), used as a fallback
            when the XML's own name doesn't contain a recognized client.

    Returns:
        str: recognized client name (one of `CLIENTES_CONHECIDOS`), or an
        empty string if neither source has one.
    """
    for fonte in (nome_cliente_xml or "", caminho_xml or ""):
        fonte_norm = fonte.upper().replace("\\", "/")
        for cliente in CLIENTES_CONHECIDOS:
            if cliente in fonte_norm:
                return cliente
    return ""


def contexto_db(tag: str):
    """Looks up application and system in the instrument registry by TAG.

    Used only as the pre-filled value for the prompt that asks the user for
    these two fields (see PdfProcessingService._processar_xml_ft); flow
    meters are usually not in this registry, so this tends to come back empty.

    Args:
        tag: Instrument TAG to look up.

    Returns:
        tuple: (aplicacao, sistema), each as a string (empty if not found
        or on a lookup error).
    """
    try:
        from data.utils_db import buscar_instrumento_por_tag
        inst = buscar_instrumento_por_tag(tag)
        if inst:
            return inst.get("aplicacao") or "", inst.get("sistema") or ""
    except Exception:
        pass
    return "", ""


def ler_nomes_assinatura():
    """Reads the default "Elaborado por"/"Verificado por" names baked into Template_Linearizacao.xlsx.

    Used only to pre-fill the rename prompt shown to the user before
    generating the report (see PdfProcessingService._processar_xml_ft) —
    the template ships with a fixed pair of names that most reports reuse
    as-is.

    Returns:
        tuple: (elaborado, verificado), each as a string (empty if the
        corresponding template cell has no value).
    """
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    caminho_template = os.path.join(base_dir, "templates", "Template_Linearizacao.xlsx")
    wb = openpyxl.load_workbook(caminho_template, read_only=True)
    try:
        ws = wb["Linearização"]
        return ws[CELLS["elaborado_nome"]].value or "", ws[CELLS["verificado_nome"]].value or ""
    finally:
        wb.close()


def _formato_decimal(casas: int) -> str:
    """Builds an Excel number format string with a fixed number of decimal places.

    Args:
        casas: number of decimal places (0 means an integer format).

    Returns:
        str: Excel number format, e.g. "0.000" for `casas=3`, or "0" for `casas=0`.
    """
    return f"0.{'0' * casas}" if casas > 0 else "0"


def _casas_significativas(valor: float, digitos: int = 7) -> int:
    """Computes how many decimal places make `valor` display with `digitos` significant figures.

    Args:
        valor: value that will be rounded/displayed.
        digitos: number of significant figures to target (default: 7).

    Returns:
        int: number of decimal places (never negative). Falls back to
        `digitos - 1` when `valor` is zero/None (order of magnitude undefined).
    """
    if not valor:
        return digitos - 1

    ordem = math.floor(math.log10(abs(valor)))
    return max(0, digitos - 1 - ordem)


def _formato_significativos(valor: float, digitos: int = 7) -> str:
    """Builds an Excel number format string that always displays `digitos` significant figures.

    Unlike `_formato_decimal` (fixed decimal places, used when the output
    must mirror the certificate's own precision), this pads with trailing
    zeros when the source value naturally has fewer significant digits —
    required for K-Factor, which must always show 7 digits regardless of
    its order of magnitude.

    Args:
        valor: value that will be written to the cell.
        digitos: number of significant figures to display (default: 7).

    Returns:
        str: Excel number format, e.g. "0.000000" for a value between 1 and
        9.999999, or "0.00000000" for a value between 0.01 and 0.09999999.
    """
    return _formato_decimal(_casas_significativas(valor, digitos))


def _celula(col: str, linha: int) -> str:
    """Builds a cell address from column and row (e.g. "C", 22 -> "C22")."""
    return f"{col}{linha}"


def _ajustar_linhas_tabela(ws, linha_ini, linha_fim_max, n_pontos):
    """Shows only the table rows actually used and normalizes their borders.

    Displays exactly the `n_pontos` used rows (from `linha_ini` to
    `linha_ini + n_pontos - 1`), hiding the rest even if they were
    originally visible, normalizes each row's border against row
    `_LINHA_BORDA_REF` (a "middle" row of the table, with a thin border) and
    closes the last used row with a thick border at the bottom.

    Args:
        ws: openpyxl worksheet.
        linha_ini: First row of the calibration table.
        linha_fim_max: Last possible row of the table (maximum capacity).
        n_pontos: Number of calibration points for this certificate.

    Notes:
        Without this normalization, re-shown rows (originally hidden, when
        there are fewer than 10 points) come out with a thicker outline, and
        the old last row (31), once it stops being the last one due to
        extra points, would be left with the thick border "leftover" in the
        middle of the table. The template's row HEIGHT is also inconsistent
        between blocks (some rows are None/default, others a fixed 15.0,
        15.75 or 14.45) — normalized here too, or a re-shown/last row comes
        out visibly taller/shorter than the rest even with a correct border.
        The "Dados a Serem Configurados" table mirrors, on row N +
        `TABELA2_OFFSET` (34), row N of this table (e.g. row 22 -> row 56) —
        visibility, border AND height normalization below are all applied to
        both tables; missing the second one leaves its old 10-point closing
        border stuck in the middle when there are more points, plus stray
        thick verticals and inconsistent heights on the re-shown rows.

        The mirror table's Frequência column (K) has the same kind of
        template inconsistency: rows 56-65 (points 1-10) carry the "0"
        format (no decimals, matching column E on table 1), but rows 66-75
        (points 11-20) come as "General" — which displays the raw,
        un-rounded value (e.g. "677,7777778") instead of the rounded
        display ("678") table 1 shows for the same point. Synced below on
        every row regardless of `usada`, since it's cheap and this is a
        display-only fix (no border/height implication).
    """
    linha_fim_usada = linha_ini + n_pontos - 1 if n_pontos else linha_ini - 1

    for row in range(linha_ini, linha_fim_max + 1):
        usada = row <= linha_fim_usada
        ws.row_dimensions[row].hidden = not usada
        ws.row_dimensions[row + TABELA2_OFFSET].hidden = not usada
        ws[f"K{row + TABELA2_OFFSET}"].number_format = ws[f"E{row}"].number_format

        if not usada:
            # Rows 22-31 of the template ship with hardcoded SAMPLE
            # calibration values (not blank) — for a certificate with fewer
            # than 10 points, these rows are hidden but their stale values
            # were otherwise still being read by KF médio (L76, AVERAGE
            # over L56:M75) and by the tolerance bounds (J79/M79, SMALL/
            # LARGE over I56:J75) on the "Dados a Serem Configurados"
            # table, since both formulas span the full 20-row range
            # regardless of how many rows are actually shown — confirmed via
            # a real Excel calculation: KF médio and the range limits came
            # out visibly wrong (e.g. upper limit = 2500, the template's own
            # sample flow rate, when the real last point was 350).
            for col in (
                CELLS["col_vazao"], CELLS["col_vol_ref"], CELLS["col_vol_med"],
                CELLS["col_mf"], CELLS["col_erro"], CELLS["col_incerteza"],
            ):
                ws[f"{col}{row}"] = None
            continue

        if row == linha_ini:
            continue  # table header row: keeps the original border

        # The original template's row height is also inconsistent between
        # blocks (some end up with default/None height, others with a fixed
        # value like 15.0/15.75/14.45) — without normalizing this, a
        # re-shown row or the last used row would come out with a different
        # height than the rest, even with the border already normalized.
        ws.row_dimensions[row].height = ws.row_dimensions[_LINHA_BORDA_REF].height
        ws.row_dimensions[row + TABELA2_OFFSET].height = ws.row_dimensions[
            _LINHA_BORDA_REF + TABELA2_OFFSET
        ].height

        for col in _BORDA_COLS:
            ws[f"{col}{row}"].border = copy(ws[f"{col}{_LINHA_BORDA_REF}"].border)
        for col in _BORDA_COLS_TABELA2:
            ws[f"{col}{row + TABELA2_OFFSET}"].border = copy(
                ws[f"{col}{_LINHA_BORDA_REF + TABELA2_OFFSET}"].border
            )

    if n_pontos:
        for col in _BORDA_COLS:
            atual = ws[f"{col}{linha_fim_usada}"].border
            ws[f"{col}{linha_fim_usada}"].border = Border(
                top=atual.top, left=atual.left, right=atual.right,
                bottom=Side(style="medium"),
            )
        for col in _BORDA_COLS_TABELA2:
            row_fim2 = linha_fim_usada + TABELA2_OFFSET
            atual = ws[f"{col}{row_fim2}"].border
            ws[f"{col}{row_fim2}"].border = Border(
                top=atual.top, left=atual.left, right=atual.right,
                bottom=Side(style="medium"),
            )


def _normalizar_linhas_tabela_generico(ws, linha_ini, linha_fim_max, n_pontos, linha_borda_ref, colunas_borda):
    """Same show/hide + border/height normalization as `_ajustar_linhas_tabela`, generalized to a single table (no 2nd-table mirroring).

    Reused for the "Falha Presumida" sheet's two tables (the calibration
    mirror, rows 22-41, and "Cálculo Falha Presumida", rows 87-106) — both
    have the exact same template bug as Linearização's own tables: only the
    first 10 rows come visible with a "normal" border, the next 10 come
    hidden with a leftover thick left/right border. Row hidden-state and
    borders are per-sheet in Excel, so hiding rows on "Linearização" doesn't
    hide the mirrored rows here — each sheet needs its own pass.

    Args:
        ws: openpyxl worksheet ("Falha Presumida").
        linha_ini: First row of the table.
        linha_fim_max: Last possible row of the table (maximum capacity).
        n_pontos: Number of calibration points for this certificate.
        linha_borda_ref: A "middle" row of the table with a normal (thin) border.
        colunas_borda: Columns to normalize the border on.
    """
    linha_fim_usada = linha_ini + n_pontos - 1 if n_pontos else linha_ini - 1

    for row in range(linha_ini, linha_fim_max + 1):
        usada = row <= linha_fim_usada
        ws.row_dimensions[row].hidden = not usada

        if not usada or row == linha_ini:
            continue  # table header row: keeps the original border

        ws.row_dimensions[row].height = ws.row_dimensions[linha_borda_ref].height
        for col in colunas_borda:
            ws[f"{col}{row}"].border = copy(ws[f"{col}{linha_borda_ref}"].border)

    if n_pontos:
        for col in colunas_borda:
            atual = ws[f"{col}{linha_fim_usada}"].border
            ws[f"{col}{linha_fim_usada}"].border = Border(
                top=atual.top, left=atual.left, right=atual.right,
                bottom=Side(style="medium"),
            )


def parear_pontos_anterior(pontos_atual: list, pontos_anterior: list) -> list:
    """Pairs each current-certificate point with the full previous-certificate point at the nearest flow rate.

    Confirmed business rule (originally for Falha Presumida, now reused by
    the primary meter AC too — see form/utils_print_ac_primario.py): these
    reports compare the same meter across two calibration events, not by
    matching table position/index — when the calibrated range differs
    between the two certificates, the nearest flow point is used (no
    interpolation, and no validation that the two certificates share the
    same meter serial number — a spare/reserve meter with a different
    serial number can legitimately be installed at the same measurement
    point).

    Args:
        pontos_atual: current certificate's points (`dados["pontos"]`), in
            the same order they were written to the Linearização table.
        pontos_anterior: previous certificate's points (`dados["pontos"]`
            from `extrair_dados_ft` on the previous XML).

    Returns:
        list[dict]: one previous-certificate point per `pontos_atual`
        entry (same shape as `extrair_dados_ft`'s "pontos" — vazao,
        meter_factor, erro_pct, incerteza, repetibilidade — so callers can
        read whichever fields they need), from the `pontos_anterior` entry
        with the closest "vazao" value.

    Raises:
        ValueError: if `pontos_anterior` is empty (nothing to pair against).
    """
    if not pontos_anterior:
        raise ValueError(
            "O certificado da calibração anterior não tem pontos de "
            "calibração para comparar."
        )

    resultado = []
    for p_atual in pontos_atual:
        vazao_atual = p_atual["vazao"]
        mais_proximo = min(
            pontos_anterior,
            key=lambda p: abs(p["vazao"] - vazao_atual),
        )
        resultado.append(mais_proximo)
    return resultado


def _parear_mf_anterior(pontos_atual: list, pontos_anterior: list) -> list:
    """Thin wrapper over `parear_pontos_anterior` returning just the meter factor — kept for `gerar_falha_presumida`.

    Returns:
        list[float]: one meter factor per `pontos_atual` entry.

    Raises:
        ValueError: if `pontos_anterior` is empty (nothing to pair against).
    """
    return [p["meter_factor"] for p in parear_pontos_anterior(pontos_atual, pontos_anterior)]


def gerar_falha_presumida(caminho_xlsx: str, dados_atual: dict, dados_anterior: dict) -> str:
    """Fills the "Falha Presumida" sheet of an already-generated Linearização workbook and exports it to its own PDF.

    Only two things are written here: the previous certificate's number
    (`FP_CERT_ANTERIOR`) and its meter factor per point (`FP_COL_MF_ANTERIOR`,
    paired by nearest flow rate — see `_parear_mf_anterior`). Everything else
    on this sheet (header, calibration table, Diff MF, Fator de Correção,
    Status, signatures) is already a template formula mirroring the
    "Linearização" sheet of the same workbook, filled in by
    `gerar_linearizacao` — see the CELLS map comment and the real cell
    mapping in tasks/TAREFAS_linearizacao_falha_presumida.md.

    Args:
        caminho_xlsx: path of the workbook already produced by
            `gerar_linearizacao` for the current certificate (same file is
            reopened and updated in place, since both sheets live in it).
        dados_atual: current certificate's data (same dict passed to
            `gerar_linearizacao` — `dados_atual["pontos"]` gives the row
            order/count already written to the Linearização table).
        dados_anterior: previous certificate's data, from `extrair_dados_ft`
            on the previous calibration's XML (its own AS_LEFT-with-
            AS_FOUND-fallback extraction already matches the confirmed rule
            for the "previous" side of this comparison).

    Returns:
        str: path of the generated "Falha Presumida" PDF (the XLSX is the
        same file passed in `caminho_xlsx`, now updated with both sheets).

    Raises:
        ValueError: if the previous certificate has no calibration points.
        PermissionError: if the output PDF already exists and is open.
    """
    wb = openpyxl.load_workbook(caminho_xlsx)
    ws = wb[SHEET_FALHA_PRESUMIDA]

    pontos_atual = dados_atual.get("pontos", [])
    mf_anteriores = _parear_mf_anterior(pontos_atual, dados_anterior.get("pontos", []))

    ws[FP_CERT_ANTERIOR] = dados_anterior.get("numero_certificado", "")
    for i, mf in enumerate(mf_anteriores):
        ws[f"{FP_COL_MF_ANTERIOR}{FP_LINHA_INI + i}"] = mf

    # Rows 87-96 of the template ship with hardcoded SAMPLE "MF Calibração
    # anterior" values (not blank) — same issue as the Linearização table
    # (see _ajustar_linhas_tabela): for a certificate with fewer than 10
    # points, the unwritten rows would keep a stale previous-MF value
    # sitting next to this run's current MF, producing a bogus Diff MF/
    # Status for a "phantom" point. Hidden, so it doesn't reach the printed
    # report, but clearing it keeps the saved XLSX free of leftover data.
    for row in range(FP_LINHA_INI + len(mf_anteriores), FP_LINHA_FIM_MAX + 1):
        ws[f"{FP_COL_MF_ANTERIOR}{row}"] = None

    _normalizar_linhas_tabela_generico(
        ws, FP_MIRROR_LINHA_INI, FP_MIRROR_LINHA_FIM_MAX, len(pontos_atual),
        FP_MIRROR_LINHA_BORDA_REF, FP_MIRROR_BORDA_COLS,
    )
    _normalizar_linhas_tabela_generico(
        ws, FP_LINHA_INI, FP_LINHA_FIM_MAX, len(pontos_atual),
        FP_LINHA_BORDA_REF_CALCULO, FP_BORDA_COLS_CALCULO,
    )

    wb.save(caminho_xlsx)

    # ── Export PDF via Excel COM ────────────────────────────────────────
    cert_limpo  = dados_atual.get("numero_certificado", "").replace(" ", "")
    tag_limpa   = dados_atual.get("tag", "").replace(" ", "")
    pasta_saida = os.path.dirname(os.path.abspath(caminho_xlsx))
    nome_pdf    = f"{cert_limpo}_{tag_limpa}_FALHA_PRESUMIDA.pdf"
    caminho_pdf = os.path.join(pasta_saida, nome_pdf)

    if os.path.exists(caminho_pdf):
        try:
            os.remove(caminho_pdf)
        except PermissionError:
            raise PermissionError(
                f"O arquivo PDF está aberto e não pode ser sobrescrito:\n{caminho_pdf}"
            )

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible        = False
    excel.DisplayAlerts  = False
    excel.ScreenUpdating = False
    excel.Interactive    = False

    try:
        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_xlsx))
        excel.Calculate()
        ws_excel = wb_excel.Worksheets(SHEET_FALHA_PRESUMIDA)
        ws_excel.ExportAsFixedFormat(0, caminho_pdf)
        wb_excel.Close(SaveChanges=False)
    finally:
        excel.Quit()

    return caminho_pdf


def gerar_linearizacao(dados: dict, caminho_xml: str) -> str:
    """Fills Template_Linearizacao.xlsx with the flow meter data and exports XLSX + PDF.

    The output files are saved in the same folder as the input XML.

    Args:
        dados: Flow meter and calibration data (header and list of points
            under "pontos").
        caminho_xml: Path of the input XML; used only to determine the
            output folder and, from the path, the client.

    Returns:
        str: Path of the generated XLSX (the PDF is generated alongside it,
        with the same base name).

    Raises:
        ValueError: If the certificate has more calibration points than
            the template supports.
        PermissionError: If the output PDF already exists and is open (it
            cannot be overwritten).
    """
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    caminho_template = os.path.join(base_dir, "templates", "Template_Linearizacao.xlsx")

    wb = openpyxl.load_workbook(caminho_template)
    ws = wb["Linearização"]

    C = CELLS  # shortcut

    # ── Left header ────────────────────────────────────────────────
    # aplicacao/sistema come from the prompt asked to the user when dropping
    # the XML (PdfProcessingService._processar_xml_ft) — they don't exist in
    # the instrument registry for a flow meter. cliente: prefers whatever
    # PdfProcessingService already resolved (auto-detected or typed in by
    # the user when detection failed — see identificar_cliente); falls back
    # to detecting it here too, so gerar_linearizacao still works standalone
    # (e.g. in tests) without going through that prompt flow.
    cliente = dados.get("cliente") or identificar_cliente(
        dados.get("cliente_xml", ""), caminho_xml
    )

    ws[C["cliente"]]         = cliente
    ws[C["instalacao"]]      = dados.get("unidade_operacional", "")
    ws[C["tag_sistema"]]     = dados.get("tag", "")
    ws[C["aplicacao"]]       = dados.get("aplicacao", "")
    ws[C["sistema"]]         = dados.get("sistema", "")
    ws[C["data_calibracao"]] = _data_br(dados.get("data_calibracao", ""))
    ws[C["tipo_medidor"]]    = dados.get("tipo", "")

    # ── Right header ─────────────────────────────────────────────────
    ws[C["num_certificado"]] = dados.get("numero_certificado", "")
    ws[C["modelo"]]          = dados.get("modelo", "")
    ws[C["fabricante"]]      = dados.get("fabricante", "")
    ws[C["num_serie"]]       = dados.get("num_serie", "")
    ws[C["diametro"]]        = dados.get("diametro", "")
    ws[C["faixa_calibrada"]] = dados.get("faixa_calibrada", "")

    # ── Nominal K-Factor ──────────────────────────────────────────────────
    # Always 7 significant figures (padded with zeros if the certificate has
    # fewer) — K-Factor doesn't have a fixed decimal place like the other
    # fields, what matters is significance, not decimal position.
    fator_k = dados.get("fator_k", 0.0)
    ws[C["fator_k"]] = fator_k
    ws[C["fator_k"]].number_format = _formato_significativos(fator_k)

    # ── Column titles (certificate's actual unit, no conversion) ────
    ws[C["titulo_vazao"]]   = f"Vazão da Calibração ({dados.get('vazao_unidade') or 'm³/h'})"
    ws[C["titulo_vol_ref"]] = f"Volume de Referência ({dados.get('vol_referencia_unidade') or 'L'})"
    ws[C["titulo_vol_med"]] = f"Volume do Medidor ({dados.get('vol_medidor_unidade') or 'L'})"

    # ── Calibration table ────────────────────────────────────────────
    # Frequência (E) and Status (S) are template formulas — calculated from
    # what we write here, do not touch. K-Factor Corrigido (O) is also a
    # formula, but has its ROUND()/format rewritten per row further below
    # (see the comment next to where column O is written).
    pontos = dados.get("pontos", [])
    linha_ini = C["tabela_linha_ini"]
    max_pontos = C["tabela_linha_fim"] - linha_ini + 1
    if len(pontos) > max_pontos:
        raise ValueError(
            f"Template de linearização suporta no máximo {max_pontos} pontos "
            f"de calibração (certificado tem {len(pontos)})."
        )

    kfc_estimados = []
    for i, p in enumerate(pontos):
        row = linha_ini + i

        cel_vazao = ws[_celula(C["col_vazao"], row)]
        cel_vazao.value = p["vazao"]
        cel_vazao.number_format = _formato_decimal(p.get("vazao_casas", 0))

        cel_vol_ref = ws[_celula(C["col_vol_ref"], row)]
        cel_vol_ref.value = p["vol_referencia"]
        cel_vol_ref.number_format = _formato_decimal(p.get("vol_referencia_casas", 0))

        cel_vol_med = ws[_celula(C["col_vol_med"], row)]
        cel_vol_med.value = p["vol_medidor"]
        cel_vol_med.number_format = _formato_decimal(p.get("vol_medidor_casas", 0))

        ws[_celula(C["col_mf"],        row)] = p["meter_factor"]
        ws[_celula(C["col_erro"],      row)] = p["erro_pct"]
        ws[_celula(C["col_incerteza"], row)] = p["incerteza"]

        # K-Factor Corrigido (O) = fator_k / meter_factor — the same operands
        # as the template formula, just recalculated here to know how many
        # decimal places to round to (the 7-significant-figure rule). Kept as
        # a formula (not a static value) so it keeps recalculating on its own
        # if the user manually edits D19/K{row} in Excel afterward.
        meter_factor = p["meter_factor"]
        try:
            kfc_estimado = fator_k / meter_factor
        except ZeroDivisionError:
            kfc_estimado = 0
        casas_kfc = _casas_significativas(kfc_estimado)
        col_kfc = C["col_kfc"]
        col_mf = C["col_mf"]

        cel_kfc = ws[_celula(col_kfc, row)]
        cel_kfc.value = f'=IFERROR(ROUND($D$19/{col_mf}{row},{casas_kfc}),"")'
        cel_kfc.number_format = _formato_decimal(casas_kfc)
        kfc_estimados.append(round(kfc_estimado, casas_kfc))

        # Mirror in the "Dados a Serem Configurados" table (col. L = "=IF(O.."):
        # same formula/value, just needs the format to match the one above.
        ws[f"L{row + TABELA2_OFFSET}"].number_format = _formato_decimal(casas_kfc)

    # KF médio (L76, template formula "=AVERAGE(...)") had a fixed format
    # with 5 decimal places — for a K-Factor of larger order of magnitude
    # (e.g. Coriolis ~20000 pulses/m³) this blew way past 7 digits (e.g.
    # "19970,57000"). Only the format is adjusted here (the value keeps the
    # template formula) — the Python average uses the same K-Factor Corrigido
    # values already written in column O, so it has the same order of
    # magnitude as the real result.
    if kfc_estimados:
        media_kfc = sum(kfc_estimados) / len(kfc_estimados)
        ws[C["kf_medio"]].number_format = _formato_significativos(media_kfc)

    # The template only ships with the first 10 rows of each table visible
    # and with a "normal" border (rows 22-31 and their mirror 56-65 in
    # "Dados a Serem Configurados"); the next 10 (32-41 / 66-75) are hidden
    # by default — shows/hides them according to this certificate's number
    # of points and normalizes the border (otherwise the re-shown rows come
    # out with a thicker outline, inherited from their style while hidden).
    _ajustar_linhas_tabela(ws, C["tabela_linha_ini"], C["tabela_linha_fim"], len(pontos))

    # ── Signatures (Elaborado por / Verificado por) ───────────────────────
    # Date always updated to today (the template ships with a fixed sample
    # date); the "Verificado por" date (N98) is already the formula "=F98"
    # in the template itself, so it follows on its own — do not write to it.
    # The names are only overwritten if the user chose to change them in the
    # prompt (PdfProcessingService._processar_xml_ft); otherwise they keep
    # the default already filled in the template (see ler_nomes_assinatura).
    ws[C["elaborado_data"]] = datetime.now()
    if dados.get("elaborado_por"):
        ws[C["elaborado_nome"]] = dados["elaborado_por"]
    if dados.get("verificado_por"):
        ws[C["verificado_nome"]] = dados["verificado_por"]

    # ── Save output XLSX ──────────────────────────────────────────────
    pasta_saida   = os.path.dirname(os.path.abspath(caminho_xml))
    cert_limpo    = dados.get("numero_certificado", "").replace(" ", "")
    tag_limpa     = dados.get("tag", "").replace(" ", "")
    nome_xlsx     = f"{cert_limpo}_{tag_limpa}_LINEARIZACAO.xlsx"
    caminho_xlsx  = os.path.join(pasta_saida, nome_xlsx)

    wb.save(caminho_xlsx)

    # ── Export PDF via Excel COM ────────────────────────────────────────
    nome_pdf    = f"{cert_limpo}_{tag_limpa}_LINEARIZACAO.pdf"
    caminho_pdf = os.path.join(pasta_saida, nome_pdf)

    if os.path.exists(caminho_pdf):
        try:
            os.remove(caminho_pdf)
        except PermissionError:
            raise PermissionError(
                f"O arquivo PDF está aberto e não pode ser sobrescrito:\n{caminho_pdf}"
            )

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible        = False
    excel.DisplayAlerts  = False
    excel.ScreenUpdating = False
    excel.Interactive    = False

    try:
        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_xlsx))
        # Frequência/KFc/Status/2nd table/KF médio/alarms are formulas —
        # forces a recalculation before exporting so they don't come out
        # blank in the PDF.
        excel.Calculate()
        # Exports only the "Linearização" sheet — ExportAsFixedFormat on the
        # Workbook (instead of the Worksheet) would also include "Falha
        # Presumida", which at this point is still unfilled (that sheet is
        # only written by gerar_falha_presumida, called separately and
        # later, if the user opts into that report).
        ws_excel = wb_excel.Worksheets("Linearização")
        ws_excel.ExportAsFixedFormat(0, caminho_pdf)
        wb_excel.Close(SaveChanges=False)
    finally:
        excel.Quit()

    return caminho_xlsx