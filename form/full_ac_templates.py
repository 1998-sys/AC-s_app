import openpyxl
from openpyxl.styles import Alignment
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
import win32com.client as win32
import os
import shutil
import tempfile
import uuid
from datetime import datetime, timedelta


def _adicionar_dia_util(data):
    data += timedelta(days=1)
    if data.weekday() == 5:  # sábado
        data += timedelta(days=2)
    elif data.weekday() == 6:  # domingo
        data += timedelta(days=1)
    return data


# Cada grupo de categorias mapeia para o mesmo título em todos os clientes desta
# família — exceto o caso "pressao" quando `permitir_split_pt_pdt` está ativo
# (hoje, só PRIO), que distingue PT de PDT pela faixa máxima calibrada.
_GRUPOS_CATEGORIA = [
    (("TERMORRESISTÊNCIA PT-100 - 2 FIOS",
      "TERMORRESISTÊNCIA PT-100 - 3 FIOS",
      "TERMORRESISTÊNCIA PT-100 - 4 FIOS"), "temperatura"),
    (("TRANSMISSOR DE TEMPERATURA COM SAÍDA EM UNIDADE ELÉTRICA",
      "TERMÔMETRO ANALÓGICO",
      "TERMÔMETRO DIGITAL"), "temperatura"),
    (("TRANSMISSOR DE PRESSÃO COM SAÍDA EM UNIDADE ELÉTRICA",
      "TRANSMISSOR DE PRESSÃO ABSOLUTA COM SAÍDA EM UNIDADE ELÉTRICA",
      "MANOMETRO DIGITAL",
      "MANOMETRO ANALÓGICO",
      "MANOMETRO DIGITAL ABSOLUTO"), "pressao"),
    (("MANOMETRO DIFERENCIAL DIGITAL",
      "MANOMETRO DIFERENCIAL ANALÓGICO"), "pressao_diferencial"),
]

_TITULOS_CATEGORIA = {
    "temperatura": "Análise Crítica de Calibração dos Sensores de Temperatura",
    "pressao": "Análise Crítica de Calibração dos Transmissores de Pressão",
    "pressao_diferencial": "Análise Crítica de Calibração dos Transmissores de Pressão Diferencial",
}


def _titulo_categoria(categoria, permitir_split_pt_pdt, max_range):
    for categorias, grupo in _GRUPOS_CATEGORIA:
        if categoria not in categorias:
            continue

        if (
            permitir_split_pt_pdt
            and grupo == "pressao"
            and categoria == "TRANSMISSOR DE PRESSÃO COM SAÍDA EM UNIDADE ELÉTRICA"
        ):
            try:
                max_range_float = float(max_range)
            except (TypeError, ValueError):
                max_range_float = None
            if max_range_float is not None and max_range_float <= 250:
                return _TITULOS_CATEGORIA["pressao_diferencial"]

        return _TITULOS_CATEGORIA[grupo]

    return None


def gerar_ac_completo(config, dados, caminho_pdf_original):
    """Preenche o template AC completo (PRIO/YINSON/YINSON ATLANTA) e exporta
    como PDF via Excel COM.

    Trata título por categoria de instrumento, quebra de linha do campo local,
    data de entrega ajustada para dia útil (+1 dia) e observações em rich text
    (range/SN atualizado ou sem alteração). Retorna o caminho absoluto do PDF.
    """
    caminho_template = config["template"]
    caminho_temp = os.path.join(
        tempfile.gettempdir(),
        f"{config['prefixo_temp']}{uuid.uuid4().hex}.xlsx"
    )
    shutil.copy(caminho_template, caminho_temp)

    wb = openpyxl.load_workbook(caminho_temp)
    ws = wb["Template Formulário"]

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "portrait"
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3

    ws["F7"] = dados.get("tag")
    ws["F8"] = dados.get("certificado")
    ws["C9"] = dados.get("data")

    if config["sistema_cell"]:
        ws[config["sistema_cell"]] = dados.get("sistema")

    # Tratamento do campo Local
    local = (dados.get("local") or "").strip()

    if len(local) > 28:
        i = local.find(" ", 28)
        if i != -1:
            local = local[:i] + "\n" + local[i + 1:]

    ws[config["local_cell"]] = local
    ws[config["local_cell"]].alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["C"].width = 22
    ws.row_dimensions[7].height = 15 * max(1, local.count("\n") + 1)

    categoria = (dados.get("categoria") or "").upper()
    titulo = _titulo_categoria(categoria, config["permitir_split_pt_pdt"], dados.get("max_range"))
    if titulo:
        ws[config["titulo_cell"]] = titulo

    # Report Date (+1 dia) - útil
    if dados.get("report_date"):
        dt = datetime.strptime(dados["report_date"], "%d/%m/%Y")
        dt_util = _adicionar_dia_util(dt)
        ws["H40"] = dt_util.strftime("%d/%m/%Y")

    range_atualizado = dados.get("range_atualizado", False)
    ns_atualizado = dados.get("sn_atualizado", False)

    blocos = []

    if range_atualizado:
        blocos.append(
            TextBlock(
                text="( X ) Sim (  ) Não\nOBSERVAÇÕES:\n",
                font=InlineFont()
            )
        )
        blocos.append(
            TextBlock(
                text="Novo range e alarmes alterados no computador de vazão\n",
                font=InlineFont(b=True)
            )
        )

    elif ns_atualizado:
        blocos.append(
            TextBlock(
                text="( X ) Sim (  ) Não\nOBSERVAÇÕES:\n",
                font=InlineFont()
            )
        )
        blocos.append(
            TextBlock(
                text="Novo NS alterado no computador de vazão / XML / SFP\n",
                font=InlineFont(b=True)
            )
        )

    else:
        blocos.append(
            TextBlock(
                text="(  ) Sim ( X ) Não\nOBSERVAÇÕES:\n",
                font=InlineFont()
            )
        )

    rich = CellRichText(*blocos)

    ws["B35"].value = rich
    ws["B35"].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(caminho_temp)
    wb.close()

    pasta_saida = os.path.dirname(os.path.abspath(caminho_pdf_original))
    certificado = dados.get("certificado", "").replace(" ", "")
    tag_limpa = dados.get("tag", "").replace(" ", "")
    nome_pdf = f"{certificado}_{tag_limpa}_AC.pdf"
    caminho_pdf_final = os.path.join(pasta_saida, nome_pdf)
    if os.path.exists(caminho_pdf_final):
        try:
            os.remove(caminho_pdf_final)
        except PermissionError:
            raise PermissionError(
                f"⚠ O arquivo PDF está aberto e não pode ser sobrescrito:\n{caminho_pdf_final}"
            )

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False
    excel.Interactive = False

    try:
        wb_excel = excel.Workbooks.Open(os.path.abspath(caminho_temp))
        wb_excel.ExportAsFixedFormat(0, caminho_pdf_final)
        wb_excel.Close(SaveChanges=False)

    finally:
        excel.Quit()
        if os.path.exists(caminho_temp):
            try:
                os.remove(caminho_temp)
            except Exception:
                pass

    return caminho_pdf_final


CONFIG_PRIO = {
    "template": "TemplateAC_PRIO.xlsx",
    "prefixo_temp": "temp_ac_prio_",
    "titulo_cell": "B2",
    "local_cell": "C7",
    "sistema_cell": "C8",
    "permitir_split_pt_pdt": True,
}

CONFIG_YINSON = {
    "template": "TemplateAC_YINSON.xlsx",
    "prefixo_temp": "temp_ac_yinson_",
    "titulo_cell": "C2",
    "local_cell": "C8",
    "sistema_cell": None,
    "permitir_split_pt_pdt": False,
}

CONFIG_YINSON_ATLANTA = {
    "template": "TemplateAC_YINSON - ATLANTA.xlsx",
    "prefixo_temp": "temp_ac_yinson_atlanta_",
    "titulo_cell": "C2",
    "local_cell": "C8",
    "sistema_cell": None,
    "permitir_split_pt_pdt": False,
}


def gerar_ac_prio(dados, caminho_pdf_original):
    """Preenche o template AC PRIO e exporta como PDF via Excel COM."""
    return gerar_ac_completo(CONFIG_PRIO, dados, caminho_pdf_original)


def gerar_ac_yinson(dados, caminho_pdf_original):
    """Preenche o template AC YINSON e exporta como PDF via Excel COM."""
    return gerar_ac_completo(CONFIG_YINSON, dados, caminho_pdf_original)


def gerar_ac_yinson_atlanta(dados, caminho_pdf_original):
    """Preenche o template AC YINSON ATLANTA e exporta como PDF via Excel COM."""
    return gerar_ac_completo(CONFIG_YINSON_ATLANTA, dados, caminho_pdf_original)
